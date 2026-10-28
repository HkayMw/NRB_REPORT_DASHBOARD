"""One-click monthly report workbook (values only, no external references)."""

from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from nrb import config
from nrb.domain import report_config as reportcfg, weeks

THIN = Side(style="thin", color="9CA3AF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(bold=True, size=14)
BAND_FILL = PatternFill("solid", fgColor="1F4E79")
BAND_FONT = Font(bold=True, color="FFFFFF", size=11)
HEAD_FILL = PatternFill("solid", fgColor="D6E4F0")
HEAD_FONT = Font(bold=True, size=10)
TOTAL_FILL = PatternFill("solid", fgColor="FDE9D9")
TOTAL_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _style(ws, cell_range, font=None, fill=None, border=BOX, align=None):
    for row in ws[cell_range]:
        for c in row:
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if border:
                c.border = border
            if align:
                c.alignment = align


def build_monthly_workbook(
    clean: pd.DataFrame,
    year: int,
    month: int,
    cfg=None,
    planned_activities: str = "",
    compiled_by: str = "",
) -> BytesIO:
    cfg = cfg or reportcfg.DEFAULT_REPORT_CONFIG
    wk_list = weeks.month_weeks(year, month)
    month_name = date(year, month, 1).strftime("%B %Y")

    # Per-week computed values: week_no -> {(gi, bi): {"rows": [...], "total": n}}
    week_values = {}
    for wno, start, end in wk_list:
        mask = (clean["date"] >= pd.Timestamp(start)) & (
            clean["date"] <= pd.Timestamp(end)
        )
        wdf = clean[mask]
        week_values[wno] = {
            (gi, bi): reportcfg.compute_block(wdf, block)
            for gi, g in enumerate(cfg["groups"])
            for bi, block in enumerate(g["blocks"])
        }

    wb = Workbook()
    _summary_sheet(
        wb.active,
        cfg,
        wk_list,
        week_values,
        month_name,
        planned_activities,
        compiled_by,
    )
    for wno, start, end in wk_list:
        _week_sheet(
            wb.create_sheet(f"Week {wno}"), cfg, wno, start, end, week_values[wno]
        )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _summary_sheet(
    ws, cfg, wk_list, week_values, month_name, planned_activities, compiled_by
):
    ws.title = "Report Summary"
    cols = reportcfg.summary_columns(cfg)
    groups: list[tuple[str, list[tuple[str, tuple[int, int, int | str]]]]] = []
    for gtitle, label, ref in cols:
        if not groups or groups[-1][0] != gtitle:
            groups.append((gtitle, []))
        groups[-1][1].append((label, ref))

    ws.cell(1, 1, config.OFFICE_NAME).font = TITLE_FONT
    ws.cell(2, 1, "MONTH").font = HEAD_FONT
    ws.cell(2, 2, month_name).font = Font(bold=True)

    # Header bands: each group gets a DATE column + its own columns.
    col = 1
    group_spans = []
    for gtitle, gcols in groups:
        span = 1 + len(gcols)
        group_spans.append((col, span))
        ws.merge_cells(
            start_row=3, start_column=col, end_row=3, end_column=col + span - 1
        )
        c = ws.cell(3, col, gtitle)
        c.fill, c.font, c.alignment = BAND_FILL, BAND_FONT, CENTER
        ws.cell(4, col, "DATE")
        for i, (label, _) in enumerate(gcols):
            ws.cell(4, col + 1 + i, label)
        col += span
    last_col = col - 1
    _style(ws, f"A3:{get_column_letter(last_col)}4", border=BOX)
    _style(
        ws,
        f"A4:{get_column_letter(last_col)}4",
        font=HEAD_FONT,
        fill=HEAD_FILL,
        align=CENTER,
    )

    # Week rows + totals.
    first_data_row = 5
    for i, (wno, start, end) in enumerate(wk_list):
        r = first_data_row + i
        col = 1
        for gtitle, gcols in groups:
            ws.cell(r, col, f"Week {wno}").font = Font(bold=True)
            for j, (label, (gi, bi, ri)) in enumerate(gcols):
                blk = week_values[wno][(gi, bi)]
                val = blk["total"] if ri == "total" else blk["rows"][ri][1]
                ws.cell(r, col + 1 + j, val).alignment = CENTER
            col += 1 + len(gcols)
    tot_row = first_data_row + len(wk_list)
    col = 1
    for gtitle, gcols in groups:
        ws.cell(tot_row, col, "TOTALS")
        for j, (label, (gi, bi, ri)) in enumerate(gcols):
            total = sum(
                (
                    week_values[w][(gi, bi)]["total"]
                    if ri == "total"
                    else week_values[w][(gi, bi)]["rows"][ri][1]
                )
                for w, _, _ in wk_list
            )
            ws.cell(tot_row, col + 1 + j, total).alignment = CENTER
        col += 1 + len(gcols)
    _style(
        ws, f"A{first_data_row}:{get_column_letter(last_col)}{tot_row - 1}", border=BOX
    )
    _style(
        ws,
        f"A{tot_row}:{get_column_letter(last_col)}{tot_row}",
        font=TOTAL_FONT,
        fill=TOTAL_FILL,
        border=BOX,
    )

    # Footer text.
    r = tot_row + 2
    if planned_activities.strip():
        ws.cell(r, 1, "PLANNED ACTIVITIES").font = Font(bold=True, underline="single")
        for line in planned_activities.strip().splitlines():
            r += 1
            ws.cell(r, 1, line)
        r += 1
    for line in compiled_by.strip().splitlines():
        r += 1
        ws.cell(r, 1, line).font = Font(bold=True)

    ws.column_dimensions["A"].width = 12
    for i in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.freeze_panes = "A5"


def _week_sheet(ws, cfg, wno, start, end, values):
    ws.cell(1, 1, config.ORG_LONG_NAME).font = Font(bold=True, size=12)
    ws.cell(2, 1, f"Week {wno}  —  {config.REPORT_TITLE}").font = TITLE_FONT
    ws.cell(3, 1, "Period:").font = HEAD_FONT
    ws.cell(3, 2, weeks.pretty_range(start, end)).font = Font(bold=True)

    r = 5
    for gi, g in enumerate(cfg["groups"]):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(r, 1, g["title"])
        c.fill, c.font, c.alignment = BAND_FILL, BAND_FONT, CENTER
        _style(ws, f"A{r}:D{r}", border=BOX)
        r += 1

        blocks = g["blocks"]
        head_row = r
        for side, block in enumerate(blocks):
            col = 1 + side * 2
            ws.merge_cells(
                start_row=head_row,
                start_column=col,
                end_row=head_row,
                end_column=col + 1,
            )
            ws.cell(head_row, col, block["title"])
        _style(
            ws,
            f"A{head_row}:D{head_row}",
            font=HEAD_FONT,
            fill=HEAD_FILL,
            align=CENTER,
            border=BOX,
        )
        r += 1

        n_rows = max(len(b.get("rows", [])) for b in blocks)
        for i in range(n_rows):
            for side, block in enumerate(blocks):
                col = 1 + side * 2
                rows = values[(gi, side)]["rows"]
                if i < len(rows):
                    label, val = rows[i]
                    ws.cell(r + i, col, label).alignment = LEFT
                    ws.cell(r + i, col + 1, val).alignment = CENTER
        _style(ws, f"A{r}:D{r + n_rows - 1}", border=BOX)
        r += n_rows

        if any(b.get("total_label") for b in blocks):
            for side, block in enumerate(blocks):
                col = 1 + side * 2
                if block.get("total_label"):
                    ws.cell(r, col, block["total_label"]).alignment = LEFT
                    ws.cell(r, col + 1, values[(gi, side)]["total"]).alignment = CENTER
            _style(ws, f"A{r}:D{r}", font=TOTAL_FONT, fill=TOTAL_FILL, border=BOX)
            r += 1
        r += 1  # gap between groups

    for column_letter, width in zip("ABCD", (30, 10, 30, 10)):
        ws.column_dimensions[column_letter].width = width


def suggested_filename(year: int, month: int) -> str:
    return (
        f"{date(year, month, 1).strftime('%B').upper()} MONTHLY REPORT_"
        f"{config.OFFICE_SHORT}.xlsx"
    )
