from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from nrb.infrastructure.exporters import excel


def test_monthly_exporter_builds_readable_summary_and_week_sheets():
    clean = pd.DataFrame(
        {
            "date": [pd.Timestamp("2026-10-01")],
            "Captured": [7],
        }
    )
    layout = {
        "groups": [
            {
                "title": "Identity",
                "summary_title": "Identity",
                "blocks": [
                    {
                        "title": "Captured",
                        "rows": [
                            {
                                "label": "Applications",
                                "fields": ["Captured"],
                                "summary_label": "Applications",
                            }
                        ],
                        "total_label": "Total captured",
                        "total_summary_label": "Total captured",
                    }
                ],
            }
        ]
    }

    output = excel.build_monthly_workbook(
        clean, 2026, 10, layout, "Field visits", "Compiled by NRB"
    )
    workbook = load_workbook(BytesIO(output.getvalue()), data_only=True)

    assert workbook.sheetnames[0] == "Report Summary"
    assert any(name.startswith("Week ") for name in workbook.sheetnames[1:])
    summary = workbook["Report Summary"]
    assert summary["B4"].value == "Applications"
    total_row = next(
        row
        for row in range(5, summary.max_row + 1)
        if summary.cell(row, 1).value == "TOTALS"
    )
    assert summary.cell(total_row, 2).value == 7
    assert "Field visits" in [
        summary.cell(row, 1).value for row in range(1, summary.max_row + 1)
    ]
    assert excel.suggested_filename(2026, 10) == (
        "OCTOBER MONTHLY REPORT_MZUZU NRB.xlsx"
    )
